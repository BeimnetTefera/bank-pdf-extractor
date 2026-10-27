"""Phase 6: approved-only Excel export (synthetic report, temporary data folder)."""

from dataclasses import replace

import pytest
from openpyxl import load_workbook

from src import review_service as review
from src.config_loader import load_config
from src.excel_exporter import (
    APPROVED_SHEET,
    EVIDENCE_COLUMNS,
    EVIDENCE_SHEET,
    METADATA_SHEET,
    METRIC_EVIDENCE_COLUMNS,
    METRIC_EVIDENCE_SHEET,
    ExportError,
    build_workbook,
    export_year,
    fiscal_year_label,
    report_years,
)
from src.analysis import analyse, report_end_year
from src.document_manager import store_document
from src.pdf_extractor import EXTRACTOR_VERSION, extract_candidates
from src.project_service import save_analysis
from src.ui.pages import excel_layout
from tests.conftest import make_pdf
from tests.test_persistence_and_review import (  # noqa: F401 (saved, conn are fixtures)
    MAPPINGS,
    METRICS,
    REVIEWER,
    STATEMENT,
    conn,
    primary_id,
    saved,
)


@pytest.fixture
def config(settings):
    return replace(load_config(), settings=settings)


def save_report(conn, config, bank: str, period: str, approve_assets: bool = True) -> str:
    """Save another synthetic report (column years follow the period) and optionally approve its total assets."""
    end_year = report_end_year(period)
    header = ("", ["Notes", str(end_year), str(end_year - 1)])
    pdf = make_pdf([[header if label == "" else (label, values) for label, values in STATEMENT]])
    document = store_document(pdf, "other.pdf", bank, period, config.settings)
    analysis = analyse(extract_candidates(pdf), replace(config, mappings=MAPPINGS), METRICS, period)
    ids = save_analysis(conn, document, analysis, EXTRACTOR_VERSION)
    if approve_assets:
        review.approve(conn, ids[("total_assets_etb", analysis.metric("total_assets_etb").primary.mapped.key)], REVIEWER)
    return document.id


def approved_rows(wb) -> list[dict]:
    rows = list(wb[APPROVED_SHEET].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def evidence_rows(wb) -> list[dict]:
    rows = list(wb[EVIDENCE_SHEET].iter_rows(values_only=True))
    assert list(rows[0]) == EVIDENCE_COLUMNS
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def test_sheets_columns_and_one_row_per_report(saved, conn, config):
    _, document, _, _ = saved
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER)
    wb = build_workbook(conn, config, [document.id], exported_by=REVIEWER)

    assert wb.sheetnames == [APPROVED_SHEET, METRIC_EVIDENCE_SHEET, EVIDENCE_SHEET, METADATA_SHEET]
    header = tuple(c.value for c in wb[APPROVED_SHEET][1])
    assert header == config.export_template.columns  # the Shega spreadsheet layout, in order
    assert header[:3] == ("bank", "fiscal_year", "Total Revenue")
    (row,) = approved_rows(wb)
    assert (row["bank"], row["fiscal_year"]) == ("Synthetic Bank", "2025/26")
    assert row["Total Assets"] == 1_234_567_000  # full Birr, not thousands
    assert row["Total Liabilities"] is None  # not approved: blank, never zero
    assert row["Profit Before Tax (Calculated)"] is None  # never calculated by the app


def test_only_approved_values_are_exported(saved, conn, config):
    _, document, _, _ = saved
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER)
    review.reject(conn, primary_id(saved, "profit_tax_etb"), REVIEWER, "Wrong row")
    review.correct(conn, primary_id(saved, "total_equity_etb"), REVIEWER, "510000000", "ETB", "Check")  # pending
    wb = build_workbook(conn, config, [document.id], exported_by=REVIEWER)

    (row,) = approved_rows(wb)
    assert row["Profit tax"] is None and row["Total Equity"] is None and row["Number of Employees"] is None
    assert [e["metric_id"] for e in evidence_rows(wb)] == ["total_assets_etb"]


def test_approved_value_without_a_column_is_reported(saved, conn, config):
    _, document, _, _ = saved
    review.manual_entry(conn, document.id, "total_assets_previous_year_etb", REVIEWER, "987654000", "ETB", 1, "prior")
    review.approve_correction(conn, document.id, "total_assets_previous_year_etb", REVIEWER)
    wb = build_workbook(conn, config, [document.id], exported_by=REVIEWER)

    (row,) = approved_rows(wb)
    assert 987_654_000 not in row.values()
    (evidence,) = evidence_rows(wb)
    assert (evidence["metric_id"], evidence["export_column"]) == ("total_assets_previous_year_etb", None)
    metadata = list(wb[METADATA_SHEET].iter_rows(values_only=True))
    assert any(k == "Warning" and "Total assets (previous year)" in v for k, v in metadata)


def test_evidence_traces_each_value_to_page_and_reviewer(saved, conn, config):
    _, document, _, _ = saved
    review.correct(conn, primary_id(saved, "total_assets_etb"), REVIEWER, "1234568000", "ETB", "Last digit misread")
    review.approve_correction(conn, document.id, "total_assets_etb", "Analyst B")
    review.manual_entry(conn, document.id, "employees", REVIEWER, "10627", "count", 2, "HR section")
    review.approve_correction(conn, document.id, "employees", "Analyst B")
    wb = build_workbook(conn, config, [document.id], exported_by=REVIEWER)

    evidence = {e["metric_id"]: e for e in evidence_rows(wb)}
    assets, employees = evidence["total_assets_etb"], evidence["employees"]
    assert (assets["value_origin"], assets["approved_value"], assets["printed_value"], assets["pdf_page"]) == (
        "corrected by reviewer", 1_234_568_000, "1,234,567", 1)
    assert (assets["original_label_or_sentence"], assets["source_sha256"]) == ("Total assets", document.sha256)
    assert (assets["entered_by"], assets["entry_reason"], assets["approved_by"]) == (
        REVIEWER, "Last digit misread", "Analyst B")
    assert (employees["value_origin"], employees["pdf_page"], employees["printed_value"]) == ("entered manually", 2, None)
    assert (employees["bank"], employees["export_column"]) == ("Synthetic Bank", "Number of Employees")
    assert approved_rows(wb)[0]["Number of Employees"] == 10627


def test_numbers_keep_their_precision_sign_and_format(saved, conn, config):
    _, document, _, _ = saved
    review.approve(conn, primary_id(saved, "profit_tax_etb"), REVIEWER)  # printed in brackets, kept as printed
    for metric_id, value, unit in (("roe_percent", "0.29", "fraction"), ("eps_etb", "12.34", "ETB"),
                                   ("total_assets_etb", "12345678901234567.89", "ETB")):
        review.manual_entry(conn, document.id, metric_id, REVIEWER, value, unit, 1, "test")
        review.approve_correction(conn, document.id, metric_id, REVIEWER, acknowledge_warnings=True)
    ws = build_workbook(conn, config, [document.id], exported_by=REVIEWER)[APPROVED_SHEET]
    cells = {ws.cell(1, col).value: ws.cell(2, col) for col in range(1, ws.max_column + 1)}

    tax, roe, eps = cells["Profit tax"], cells["Return on Equity (ROE) (Reported)"], cells["Earning per share (EPS) (Reported)"]
    assert (tax.value, tax.number_format) == (-4_313_629_000, "#,##0")
    assert (roe.value, roe.number_format) == (0.29, "0.00%")
    assert (eps.value, eps.number_format) == (12.34, "#,##0.00")
    # Too many digits for an Excel number: kept as exact text rather than rounded.
    assert cells["Total Assets"].value == "12345678901234567.89"


def test_nothing_approved_is_an_error(saved, conn, config):
    _, document, _, _ = saved
    with pytest.raises(ExportError, match="Nothing is approved"):
        build_workbook(conn, config, [document.id], exported_by=REVIEWER)


def test_unknown_report_is_an_error(saved, conn, config):
    with pytest.raises(ExportError, match="not in the database"):
        build_workbook(conn, config, ["missing"], exported_by=REVIEWER)


def test_templates_page_lists_every_excel_column(config):
    rows = excel_layout(config)
    assert [r["Excel column"] for r in rows] == list(config.export_template.columns)
    by_column = {r["Excel column"]: r for r in rows}
    assert by_column["Total Assets"]["Filled by"] == "Total assets"
    assert by_column["Total Assets"]["Group"] == "Extracted"
    assert by_column["Profit Before Tax (Calculated)"]["Group"] == "Calculated"
    groups = [r["Group"] for r in rows]
    assert (groups.count("Extracted"), groups.count("Calculated")) == (39, 26)
    assert all(r["Status"] == "Filled when approved" for r in rows if r["Group"] == "Extracted")


# --- one workbook per fiscal year ---------------------------------------------------------

def test_fiscal_year_label():
    assert fiscal_year_label(2026) == "2025/26" and fiscal_year_label(2000) == "1999/00"


def test_one_workbook_per_year_with_every_bank(saved, conn, config):
    _, document, _, _ = saved  # Synthetic Bank, saved as "2025/26"
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER)
    other = save_report(conn, config, "Zemen Bank", "2026")  # same fiscal year, typed differently (saved as 2025/26)
    next_year = save_report(conn, config, "Awash Bank", "2026/27")
    assert report_years(conn) == [2027, 2026]

    result = export_year(conn, config, 2026, exported_by=REVIEWER)
    assert result.path == config.settings.exports_dir / "approved_2025-26.xlsx"
    assert (result.fiscal_year, result.reports, result.values, result.warnings) == ("2025/26", 2, 2, ())
    wb = load_workbook(result.path)  # opens as a real workbook
    rows = approved_rows(wb)
    assert [(r["bank"], r["fiscal_year"]) for r in rows] == [
        ("Synthetic Bank", "2025/26"), ("Zemen Bank", "2025/26")]  # one Shega form, sorted by bank
    evidence = {e["source_document_id"]: e for e in evidence_rows(wb)}
    assert set(evidence) == {document.id, other}
    assert evidence[other]["report_year_as_saved"] == "2025/26"  # Save stores YYYY/YY (D-085)
    metadata = dict(wb[METADATA_SHEET].iter_rows(min_row=2, values_only=True))
    assert (metadata["Fiscal year"], metadata["Reports"], metadata["Exported by"]) == ("2025/26", 2, REVIEWER)

    following = export_year(conn, config, 2027, exported_by=REVIEWER)
    assert following.path.name == "approved_2026-27.xlsx"
    assert [e["source_document_id"] for e in evidence_rows(load_workbook(following.path))] == [next_year]
    assert not list(config.settings.exports_dir.glob("*.tmp"))


def test_exporting_a_year_again_replaces_its_file(saved, conn, config):
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER)
    first = export_year(conn, config, 2026, exported_by=REVIEWER)
    save_report(conn, config, "Zemen Bank", "2025/26")
    second = export_year(conn, config, 2026, exported_by=REVIEWER)
    assert second.path == first.path and second.reports == 2
    assert list(config.settings.exports_dir.glob("*.xlsx")) == [second.path]


def test_a_bank_saved_twice_in_a_year_is_flagged(saved, conn, config):
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER)
    save_report(conn, config, "Synthetic Bank", "2026")
    result = export_year(conn, config, 2026, exported_by=REVIEWER)
    (warning,) = result.warnings
    assert "Synthetic Bank has 2 saved reports" in warning
    metadata = list(load_workbook(result.path)[METADATA_SHEET].iter_rows(values_only=True))
    assert ("Warning", warning) in metadata


def test_reports_without_approvals_are_left_out(saved, conn, config):
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER)
    save_report(conn, config, "Zemen Bank", "2025/26", approve_assets=False)
    assert export_year(conn, config, 2026, exported_by=REVIEWER).reports == 1


def test_year_without_approvals_or_reports_is_an_error(saved, conn, config):
    with pytest.raises(ExportError, match="Nothing is approved yet for 2025/26"):
        export_year(conn, config, 2026, exported_by=REVIEWER)
    with pytest.raises(ExportError, match="No reports are saved for 2030/31"):
        export_year(conn, config, 2031, exported_by=REVIEWER)


# --- Evidence links (spec section 1) -------------------------------------------------------

def test_value_cells_link_to_their_evidence_page_with_a_comment(saved, conn, config):
    _, document, _, _ = saved
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER, note="Statement, page 1")
    wb = build_workbook(conn, config, [document.id], exported_by=REVIEWER)
    ws = wb[APPROVED_SHEET]
    header = [c.value for c in ws[1]]
    cell = ws.cell(2, header.index("Total Assets") + 1)
    assert cell.value == 1_234_567_000 and cell.number_format == "#,##0"  # still a number
    assert cell.hyperlink.target == f"http://localhost:8501/evidence?report={document.id}&metric=total_assets_etb"
    assert "PDF figure, PDF page 1" in cell.comment.text and "Note: Statement, page 1" in cell.comment.text
    blank = ws.cell(2, header.index("Total Liabilities") + 1)
    assert blank.value is None and blank.hyperlink is None


def test_evidence_sheet_lists_every_analysed_metric(saved, conn, config):
    _, document, _, _ = saved
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER)
    review.correct(conn, primary_id(saved, "total_equity_etb"), REVIEWER, "510000000", "ETB", "Note 30 says so")
    wb = build_workbook(conn, config, [document.id], exported_by=REVIEWER)
    rows = list(wb[METRIC_EVIDENCE_SHEET].iter_rows(values_only=True))
    assert list(rows[0]) == METRIC_EVIDENCE_COLUMNS
    by_metric = {r[2]: dict(zip(rows[0], r)) for r in rows[1:]}
    assert set(by_metric) == set(METRICS)
    assets, equity, employees = by_metric["total_assets_etb"], by_metric["total_equity_etb"], by_metric["employees"]
    assert (assets["status"], assets["value"], assets["source_type"], assets["pdf_page"]) == (
        "Approved", 1_234_567_000, "PDF figure", 1)
    assert (equity["status"], equity["value"], equity["note"]) == ("Correction pending", None, "Note 30 says so")
    assert (employees["status"], employees["value"]) == ("Not found", None)  # approved values only (D-032)
    assert assets["evidence_url"].endswith(f"report={document.id}&metric=total_assets_etb")


def test_source_evidence_has_the_evidence_address(saved, conn, config):
    _, document, _, _ = saved
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER)
    (row,) = evidence_rows(build_workbook(conn, config, [document.id], exported_by=REVIEWER))
    assert row["evidence_url"] == f"http://localhost:8501/evidence?report={document.id}&metric=total_assets_etb"


def test_app_url_from_settings_is_used(saved, conn, config):
    _, document, _, _ = saved
    review.approve(conn, primary_id(saved, "total_assets_etb"), REVIEWER)
    other = replace(config, settings=replace(config.settings, app_url="http://reviewpc:8600"))
    (row,) = evidence_rows(build_workbook(conn, other, [document.id], exported_by=REVIEWER))
    assert row["evidence_url"].startswith("http://reviewpc:8600/evidence?")
